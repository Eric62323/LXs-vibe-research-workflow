import sys
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_border(borders_node, name, val="single", sz="4", space="0", color="auto"):
    """添加边框节点"""
    element = OxmlElement(f'w:{name}')
    element.set(qn('w:val'), val)
    element.set(qn('w:sz'), sz)  # sz: 4=0.5pt, 6=0.75pt, 12=1.5pt
    element.set(qn('w:space'), space)
    element.set(qn('w:color'), color)
    borders_node.append(element)

def set_cell_bottom_border(cell, sz="4"):
    """给单元格设置底部边框"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    
    # 移除旧的 bottom
    old_bottom = tcBorders.find(qn('w:bottom'))
    if old_bottom is not None:
        tcBorders.remove(old_bottom)
        
    add_border(tcBorders, 'bottom', sz=sz)

def format_docx(file_path):
    if not os.path.exists(file_path):
        print(f"Error: File not found: {file_path}")
        return

    try:
        doc = Document(file_path)
    except Exception as e:
        print(f"Error opening file: {e}")
        return
    
    print(f"Processing tables in: {file_path}...")
    table_count = len(doc.tables)
    
    for i, table in enumerate(doc.tables):
        # 1. 设置表格整体边框 (顶底线)
        tbl = table._tbl
        tblPr = tbl.tblPr
        
        # 移除旧的 tblBorders
        tblBorders = tblPr.find(qn('w:tblBorders'))
        if tblBorders is not None:
            tblPr.remove(tblBorders)
        
        # 创建新的 tblBorders
        tblBorders = OxmlElement('w:tblBorders')
        # 0.75pt (6 units), 0.5pt (4 units)
        add_border(tblBorders, 'top', sz="6")      # 顶线 0.75pt
        add_border(tblBorders, 'bottom', sz="6")   # 底线 0.75pt
        add_border(tblBorders, 'left', val="nil")
        add_border(tblBorders, 'right', val="nil")
        add_border(tblBorders, 'insideH', val="nil") # 去除内部横线
        add_border(tblBorders, 'insideV', val="nil") # 去除内部竖线
        tblPr.append(tblBorders)

        # 2. 设置第一行下边框（栏目线）- 0.5pt
        if len(table.rows) > 0:
            for cell in table.rows[0].cells:
                set_cell_bottom_border(cell, sz="4") 

        # 3. 字体和行距设置
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    # 单倍行距
                    if paragraph.paragraph_format.line_spacing != 1.0:
                        paragraph.paragraph_format.line_spacing = 1.0
                    
                    # 确保段前段后间距为0（可选，视学术规范而定，这里为了紧凑通常设为0）
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)

                    for run in paragraph.runs:
                        # 字号：五号 (10.5pt)
                        run.font.size = Pt(10.5)
                        
                        # 中英文字体设置
                        run.font.name = 'Times New Roman' # 西文
                        if run._element.rPr is None:
                            run._element.get_or_add_rPr()
                        if run._element.rPr.rFonts is None:
                            run._element.rPr.get_or_add_rFonts()
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体') # 中文

    doc.save(file_path)
    print(f"Successfully formatted {table_count} tables in {file_path}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        format_docx(sys.argv[1])
    else:
        print("Usage: python table_formatter.py <docx_file_path>")
